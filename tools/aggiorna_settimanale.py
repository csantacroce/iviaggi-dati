#!/usr/bin/env python3
"""Aggiornamento automatico settimanale (lo lancia GitHub Actions ogni lunedì mattina).

1. Prezzi dei carburanti: scarica dal Ministero (MIMIT, Osservaprezzi Carburanti, open data IODL 2.0)
   il file dei prezzi praticati e riscrive il foglio «prezzi» dell'Excel solo per gli impianti già
   abbinati alle aree nel foglio «impianti_ministero». Gli altri fogli (e le correzioni a mano) non si toccano.
2. Cambi: scarica i tassi di riferimento della Banca Centrale Europea e riscrive il foglio «tassi_cambio».

Poi tools/build.py controlla tutto e pubblica una nuova versione solo se i dati sono cambiati.
Se una fonte non risponde o il file è strano, quella parte NON viene modificata (l'altra sì).
Se falliscono entrambe esce con errore: GitHub manda una mail e i dati pubblicati restano quelli di prima.

Uso:  python3 tools/aggiorna_settimanale.py [--prezzi FILE.csv] [--cambi FILE.xml]
      (con i file locali non scarica nulla: serve per le prove)
"""
import csv
import datetime as dt
import io
import os
import re
import sys
import urllib.request
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
XLSX = os.path.join(ROOT, "source", "dati.xlsx")
URL_PREZZI = "https://www.mimit.gov.it/images/exportCSV/prezzo_alle_8.csv"
URL_CAMBI = "https://www.ecb.europa.eu/stats/eurofxref/eurofxref-daily.xml"
VALUTE = ["USD", "GBP", "CHF", "SEK", "NOK", "DKK", "PLN", "CZK", "HUF", "RON"]   # quelle del menu Impostazioni
MIN_COPERTURA = 0.5     # se meno della metà degli impianti abbinati è nel file, il file è sospetto: non si usa


def scarica(url):
    req = urllib.request.Request(url, headers={"User-Agent": "iVIAGGI-dati/1.0 (aggiornamento settimanale)"})
    with urllib.request.urlopen(req, timeout=120) as r:
        return r.read()


# ---------------------------------------------------------------- prezzi

def leggi_prezzi(raw):
    """Legge il CSV del Ministero in modo tollerante (separatore ; oppure |, riga «Estrazione del …» in testa,
    virgola o punto decimale). Restituisce (data_estrazione 'AAAA-MM-GG' o None, righe)."""
    testo = raw.decode("utf-8-sig", errors="replace")
    righe = [r for r in testo.splitlines() if r.strip()]
    estrazione = None
    m = re.search(r"(\d{4}-\d{2}-\d{2})|(\d{2})/(\d{2})/(\d{4})", righe[0]) if righe and "estrazione" in righe[0].lower() else None
    if m:
        estrazione = m.group(1) or f"{m.group(4)}-{m.group(3)}-{m.group(2)}"
    # intestazione: la prima riga che contiene "idimpianto"
    start = next((i for i, r in enumerate(righe[:5]) if "idimpianto" in r.lower().replace("_", "")), None)
    if start is None:
        raise ValueError("intestazione del file prezzi non trovata")
    sep = max(["|", ";", ","], key=lambda s: righe[start].count(s))
    reader = csv.reader(io.StringIO("\n".join(righe[start:])), delimiter=sep)
    head = [h.strip().lower().replace("_", "") for h in next(reader)]

    def col(*nomi):
        for n in nomi:
            if n in head:
                return head.index(n)
        raise ValueError(f"colonna mancante nel file prezzi: {nomi[0]}")
    c_id, c_desc, c_prezzo = col("idimpianto"), col("desccarburante", "carburante"), col("prezzo")
    c_self, c_data = col("isself", "self"), col("dtcomu", "datacomunicazione", "data")
    out = []
    for r in reader:
        if len(r) <= max(c_id, c_desc, c_prezzo, c_self, c_data):
            continue
        try:
            out.append({
                "id": int(r[c_id].strip()),
                "desc": r[c_desc].strip(),
                "prezzo": float(r[c_prezzo].strip().replace(",", ".")),
                "self": r[c_self].strip() in ("1", "true", "True", "S", "s"),
                "data": data_iso(r[c_data].strip()),
            })
        except ValueError:
            continue
    return estrazione, out


def data_iso(s):
    """'18/10/2026 07:12:45' oppure '2026-10-18 07:12:45' → '2026-10-18 07:12'."""
    for f in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S"):
        try:
            return dt.datetime.strptime(s, f).strftime("%Y-%m-%d %H:%M")
        except ValueError:
            pass
    raise ValueError(s)


FAMIGLIE = {   # carburante dell'app: (nome esatto del Ministero, parole per riconoscere le varianti)
    "Benzina": ("benzina", ("benzina", "super")),
    "Gasolio": ("gasolio", ("gasolio", "diesel")),
    "GPL": ("gpl", ("gpl",)),
    "Metano": ("metano", ("metano", "gnc")),
}


def scegli(prezzi_impianto):
    """Per ogni carburante dell'app e modalità (self/servito): il prodotto con il nome esatto ("Gasolio");
    se manca, la variante più economica della stessa famiglia (le versioni premium costano di più).
    Esclusi GNL (gas liquido per camion) e prezzi fuori scala."""
    scelti = {}
    for app_nome, (esatto, parole) in FAMIGLIE.items():
        for self_ in (True, False):
            cand = [p for p in prezzi_impianto if p["self"] == self_ and 0.3 < p["prezzo"] < 5]
            esatti = [p for p in cand if p["desc"].lower() == esatto]
            famiglia = [p for p in cand if any(w in p["desc"].lower() for w in parole) and "gnl" not in p["desc"].lower()
                        and not (app_nome == "Benzina" and "gasolio" in p["desc"].lower())]
            pool = esatti or famiglia
            if pool:
                best = max(pool, key=lambda p: p["data"]) if esatti else min(pool, key=lambda p: p["prezzo"])
                scelti[(app_nome, "self" if self_ else "servito")] = best
    return scelti


def aggiorna_prezzi(wb, raw, oggi):
    estrazione, righe = leggi_prezzi(raw)
    ws_imp = wb["impianti_ministero"]
    h_imp = [c.value for c in ws_imp[1]]
    i_area, i_imp = h_imp.index("id_area"), h_imp.index("id_impianto")
    abbinati = {}
    for r in ws_imp.iter_rows(min_row=2, values_only=True):
        if r[i_area] is not None and r[i_imp] is not None:
            abbinati[int(r[i_imp])] = int(r[i_area])
    per_impianto = {}
    for p in righe:
        if p["id"] in abbinati:
            per_impianto.setdefault(p["id"], []).append(p)
    copertura = len(per_impianto) / max(1, len(abbinati))
    if copertura < MIN_COPERTURA:
        raise ValueError(f"nel file ci sono solo {len(per_impianto)} dei {len(abbinati)} impianti abbinati: file sospetto, prezzi non aggiornati")

    ws = wb["prezzi"]
    head = [c.value for c in ws[1]]
    pos = {h: i for i, h in enumerate(head)}
    vecchie = [list(r) for r in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r[:7])]
    # si tengono le righe degli impianti che oggi non sono nel file (prezzi non comunicati di recente)
    tenute = [r for r in vecchie if r[pos["id_impianto"]] not in per_impianto]
    nuove = []
    for imp, lista in sorted(per_impianto.items()):
        for (carb, mod), p in sorted(scegli(lista).items()):
            nuove.append({"id_area": abbinati[imp], "id_impianto": imp, "carburante": carb, "modalita": mod,
                          "prezzo": round(p["prezzo"], 3), "data_comunicazione": p["data"],
                          "data_estrazione": estrazione or oggi})
    # riscrive il foglio: stesse colonne, la colonna di aiuto _area con la formula
    ws.delete_rows(2, ws.max_row)
    r_idx = 2
    for r in tenute:
        for ci in range(7):
            ws.cell(row=r_idx, column=ci + 1, value=r[ci])
        r_idx += 1
    for n in nuove:
        for h, v in n.items():
            ws.cell(row=r_idx, column=pos[h] + 1, value=v)
        r_idx += 1
    if "_area" in pos:
        for rr in range(2, r_idx):
            ws.cell(row=rr, column=pos["_area"] + 1,
                    value=f'=IF(A{rr}="","",IFERROR(INDEX(aree_servizio!B:B,MATCH(A{rr},aree_servizio!A:A,0)),"?"))')
    imposta_info(wb, "data_estrazione", estrazione or oggi)
    return f"prezzi: {len(nuove)} righe da {len(per_impianto)}/{len(abbinati)} impianti (estrazione {estrazione or oggi}), {len(tenute)} righe vecchie tenute"


# ---------------------------------------------------------------- cambi

def leggi_cambi(raw):
    root = ET.fromstring(raw)
    data, tassi = None, {}
    for el in root.iter():
        if el.tag.endswith("Cube") and "time" in el.attrib:
            data = el.attrib["time"]
        if el.tag.endswith("Cube") and "currency" in el.attrib:
            tassi[el.attrib["currency"]] = float(el.attrib["rate"])
    mancanti = [v for v in VALUTE if v not in tassi]
    if not data or mancanti:
        raise ValueError(f"file BCE incompleto (mancano: {', '.join(mancanti) or 'data'})")
    return data, tassi


def aggiorna_cambi(wb, raw):
    data, tassi = leggi_cambi(raw)
    if "tassi_cambio" not in wb.sheetnames:
        wb.create_sheet("tassi_cambio")
    ws = wb["tassi_cambio"]
    ws.delete_rows(1, ws.max_row)
    ws.append(["valuta", "per_euro", "data", "fonte"])
    for v in VALUTE:
        ws.append([v, tassi[v], data, "BCE, tassi di riferimento (eurofxref)"])
    imposta_info(wb, "data_cambi", data)
    return f"cambi BCE del {data}: " + ", ".join(f"{v} {tassi[v]}" for v in VALUTE)


def imposta_info(wb, chiave, valore):
    ws = wb["info_aggiornamento"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == chiave:
            row[1].value = valore
            return
    ws.append([chiave, valore])


def main():
    args = sys.argv[1:]
    file_prezzi = args[args.index("--prezzi") + 1] if "--prezzi" in args else None
    file_cambi = args[args.index("--cambi") + 1] if "--cambi" in args else None
    oggi = dt.date.today().isoformat()
    wb = load_workbook(XLSX)
    esiti, errori = [], []
    try:
        raw = open(file_prezzi, "rb").read() if file_prezzi else scarica(URL_PREZZI)
        esiti.append(aggiorna_prezzi(wb, raw, oggi))
    except Exception as e:
        errori.append(f"prezzi non aggiornati: {e}")
    try:
        raw = open(file_cambi, "rb").read() if file_cambi else scarica(URL_CAMBI)
        esiti.append(aggiorna_cambi(wb, raw))
    except Exception as e:
        errori.append(f"cambi non aggiornati: {e}")
    if esiti:
        wb.save(XLSX)
    for x in esiti:
        print("OK  ", x)
    for x in errori:
        print("ATTENZIONE", x)
    if not esiti:
        sys.exit(1)


if __name__ == "__main__":
    main()
