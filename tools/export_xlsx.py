#!/usr/bin/env python3
"""Crea l'Excel da modificare a partire dal database (serve una volta, o per ripartire da zero).
Uso: python3 export_xlsx.py iviaggi_database.sqlite source/dati.xlsx"""
import os
import sqlite3
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from dbtools import load_schema

SPARE = 150  # righe vuote già pronte (con i nomi di aiuto) in fondo ai fogli dove si aggiungono righe
LOOKUPS = {  # foglio: [(intestazione aiuto, colonna sorgente, foglio di ricerca, colonna id, colonna del nome)]
    "aree_servizio": [("_autostrada", "id_autostrada", "autostrade", "id", "sigla"),
                      ("_direzione", "id_casello_direzione", "caselli", "id", "nome")],
    "aree_colonnine": [("_area", "id_area", "aree_servizio", "id", "nome"),
                       ("_gestore", "id_gestore", "gestori_ricarica", "id", "nome")],
    "aree_marche_carburante": [("_area", "id_area", "aree_servizio", "id", "nome"),
                               ("_marca", "id_marca", "marche_carburante", "id", "nome")],
    "aree_marche_ristorante": [("_area", "id_area", "aree_servizio", "id", "nome"),
                               ("_marca", "id_marca", "marche_ristorante", "id", "nome")],
    "aree_servizi": [("_area", "id_area", "aree_servizio", "id", "nome"),
                     ("_servizio", "id_servizio", "servizi", "id", "nome")],
    "prezzi": [("_area", "id_area", "aree_servizio", "id", "nome")],
    "impianti_ministero": [("_area", "id_area", "aree_servizio", "id", "nome")],
    "tariffe_ricarica": [("_gestore", "id_gestore", "gestori_ricarica", "id", "nome")],
}
SPARE_SHEETS = {"tariffe_ricarica", "aree_colonnine", "aree_marche_carburante", "aree_marche_ristorante", "aree_servizi"}

LEGGIMI = [
    "COME SI USA QUESTO FILE",
    "",
    "1. Ogni foglio è una tabella del database dell'app. Modifica solo le celle: non cambiare i nomi dei fogli né le intestazioni della prima riga.",
    "2. Le colonne che iniziano con _ (verdi) sono solo di aiuto: mostrano il nome dell'area, dell'autostrada o del gestore. Non vengono lette.",
    "3. Per aggiungere una colonnina: nel foglio aree_colonnine scrivi una nuova riga in fondo (id nuovo = ultimo id + 1; id_area = numero dell'area, lo trovi nel foglio aree_servizio; id_gestore dal foglio gestori_ricarica).",
    "4. Per nascondere un'area dall'app: nel foglio aree_servizio metti attiva_2026 = 0. Per mostrarla: 1. (-1 = esclusa, non è su autostrada).",
    "4b. Listino ricarica (foglio tariffe_ricarica): una riga per gestore, corrente (AC/DC) e fascia di potenza (potenza_min_kw–potenza_max_kw). Prezzi a consumo senza abbonamento, IVA inclusa. stato = 'verificato' oppure 'da verificare'. data_fonte = data del listino. Il prezzo scritto in una riga di aree_colonnine (es. A22) vince sul listino.",
    "5. Per segnare un'area chiusa: colonna stato = 'chiusa (temporaneamente)'; per riaprirla: 'aperta'.",
    "6. Date nel formato AAAA-MM-GG (per esempio 2026-10-14). Prezzi e potenze con il punto o la virgola come li scrive il tuo Excel.",
    "7. Celle vuote sono valide (valore mancante). Righe completamente vuote vengono ignorate.",
    "8. Salva con lo stesso nome (dati.xlsx) e caricalo dalla pagina di aggiornamento. Se c'è un errore, la pagina te lo spiega in italiano e l'app continua a usare i dati vecchi.",
    "",
    "COSA NON FARE: non cambiare gli id di righe esistenti, non cancellare un'area se ha righe collegate (colonnine, marche, servizi, prezzi): togli prima quelle righe, o metti attiva_2026 = 0.",
]


def main(db_path, out_path):
    here = os.path.dirname(os.path.abspath(__file__))
    schema = load_schema(sys.argv[3] if len(sys.argv) > 3 else os.path.join(here, "..", "source", "schema.sql"))
    con = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "LEGGIMI"
    for i, line in enumerate(LEGGIMI, 1):
        ws.cell(row=i, column=1, value=line)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 150
    head_fill = PatternFill("solid", fgColor="00693E")
    help_fill = PatternFill("solid", fgColor="DDEFE3")
    for table, cols in schema.items():
        ws = wb.create_sheet(table)
        names = [c[0] for c in cols]
        extra = LOOKUPS.get(table, [])
        header = names + [e[0] for e in extra]
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill if not str(cell.value).startswith("_") else PatternFill("solid", fgColor="5E8F6F")
            cell.alignment = Alignment(vertical="center")
        rows = con.execute(f'SELECT {",".join(chr(34)+n+chr(34) for n in names)} FROM "{table}"').fetchall()
        total = len(rows) + (SPARE if table in SPARE_SHEETS else 0)
        for r in range(2, total + 2):
            if r - 2 < len(rows):
                for ci, v in enumerate(rows[r - 2], 1):
                    if v == "":
                        v = None
                    ws.cell(row=r, column=ci, value=v)
            for k, (hdr, src, sheet, idcol, namecol) in enumerate(extra):
                s = get_column_letter(names.index(src) + 1)
                lk_id = get_column_letter([c[0] for c in schema[sheet]].index(idcol) + 1)
                lk_nm = get_column_letter([c[0] for c in schema[sheet]].index(namecol) + 1)
                cell = ws.cell(row=r, column=len(names) + k + 1,
                               value=f'=IF({s}{r}="","",IFERROR(INDEX({sheet}!{lk_nm}:{lk_nm},MATCH({s}{r},{sheet}!{lk_id}:{lk_id},0)),"?"))')
                cell.fill = help_fill
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{total + 1}"
        for ci, h in enumerate(header, 1):
            longest = max([len(str(h))] + [len(str(x[ci - 1])) for x in rows[:200] if ci <= len(names) and x[ci - 1] is not None])
            ws.column_dimensions[get_column_letter(ci)].width = min(max(10, longest + 2), 60)
    wb.save(out_path)
    print("Excel creato:", out_path)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
